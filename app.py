# -*- coding: utf-8 -*-
"""KORA AI - assistente personale in italiano (Streamlit + Gemini 2.5 Flash)."""
import re, io, time, hashlib, random
from urllib.parse import quote
from xml.sax.saxutils import escape

import streamlit as st
from google import genai
from google.genai import types
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from docx import Document

MODEL = "gemini-2.5-flash"

SYSTEM_INSTRUCTION = """Sei KORA AI, assistente personale in lingua italiana.
RUOLO: esperto professionale di elettronica, elettrotecnica (Ohm, Kirchhoff, Thevenin/Norton,
circuiti in AC/DC, trifase, trasformatori, motori, elettronica analogica e digitale,
microcontrollori, PCB, normative), informatica (Python, C/C++, web, reti, database, sicurezza),
logica, matematica e lingua italiana.
STILE: rispondi subito, in modo diretto ed esplicito, senza giri di parole né preamboli.
Prima la risposta, poi (solo se utile) formule, passaggi e unità di misura corrette.
Il codice deve essere completo, funzionante e commentato in modo professionale.
Se la richiesta è ambigua fai UNA sola domanda breve. Non inventare: se non sai, dillo.
Rispondi sempre in italiano."""

IMG_RE = re.compile(
    r"\b(crea|creami|genera|generami|disegna|disegnami|fammi|realizza)\b.*"
    r"\b(immagine|foto|disegno|illustrazione|logo|immagini)\b", re.I)

LOGO_SVG = """
<svg viewBox="0 0 200 200" width="{s}" height="{s}" xmlns="http://www.w3.org/2000/svg">
 <defs>
  <linearGradient id="g" x1="0" y1="0" x2="1" y2="1">
   <stop offset="0" stop-color="#00E5FF"/><stop offset="1" stop-color="#8A5CFF"/></linearGradient>
  <filter id="b"><feGaussianBlur stdDeviation="4"/></filter>
 </defs>
 <circle cx="100" cy="100" r="82" fill="none" stroke="url(#g)" stroke-width="10" filter="url(#b)" opacity=".7"/>
 <circle cx="100" cy="100" r="82" fill="#0B1230" stroke="url(#g)" stroke-width="5"/>
 <g stroke="url(#g)" stroke-width="3" fill="none" stroke-linecap="round" opacity=".8">
  <path d="M18 100h28M154 100h28M100 18v28M100 154v28"/>
  <circle cx="18" cy="100" r="4"/><circle cx="182" cy="100" r="4"/>
  <circle cx="100" cy="18" r="4"/><circle cx="100" cy="182" r="4"/></g>
 <path d="M72 60v80M72 100l50-40M72 100l50 40" stroke="url(#g)" stroke-width="12"
  stroke-linecap="round" stroke-linejoin="round" fill="none"/>
 <circle cx="128" cy="100" r="6" fill="#00E5FF"/>
</svg>"""

CSS = """
<style>
#MainMenu, footer, header [data-testid="stToolbar"], [data-testid="stDecoration"]{display:none!important}
.stApp{background:radial-gradient(1200px 600px at 20% -10%,#1a1f5a55,transparent),
       radial-gradient(900px 500px at 100% 0,#00e5ff22,transparent),#070B1A}
.block-container{padding:1rem .8rem 6rem;max-width:860px}
.kora-title{font-size:2rem;font-weight:800;letter-spacing:.25em;
 background:linear-gradient(90deg,#00E5FF,#8A5CFF);-webkit-background-clip:text;
 -webkit-text-fill-color:transparent;margin:0}
.kora-sub{color:#7f93c7;font-size:.8rem;letter-spacing:.2em;margin-top:-4px}
[data-testid="stChatMessage"]{background:#0F163099;border:1px solid #00e5ff22;
 border-radius:18px;backdrop-filter:blur(8px);box-shadow:0 0 18px #00e5ff10}
.stButton>button,.stDownloadButton>button{border-radius:12px;border:1px solid #00e5ff66;
 background:linear-gradient(135deg,#0F1630,#1b2260);color:#e8f1ff}
.stButton>button:hover,.stDownloadButton>button:hover{border-color:#00E5FF;box-shadow:0 0 12px #00e5ff55}
.splash{display:flex;flex-direction:column;align-items:center;justify-content:center;height:70vh;
 animation:pulse 1.4s ease-in-out infinite alternate}
@keyframes pulse{from{filter:drop-shadow(0 0 8px #00e5ff55)}to{filter:drop-shadow(0 0 30px #8a5cffaa)}}
@media(max-width:640px){.kora-title{font-size:1.5rem}}
</style>"""

st.set_page_config(page_title="KORA AI", page_icon="🤖", layout="centered",
                   initial_sidebar_state="collapsed")
st.markdown(CSS, unsafe_allow_html=True)

# ---------- Secrets / client ----------
GEMINI_KEY = st.secrets.get("GEMINI_API_KEY")
if not GEMINI_KEY:
    st.error("Manca GEMINI_API_KEY nei Secrets di Streamlit.")
    st.stop()

@st.cache_resource
def get_client():
    return genai.Client(api_key=GEMINI_KEY)
client = get_client()

# ---------- Stato ----------
ss = st.session_state
ss.setdefault("messages", [])
ss.setdefault("splash_done", False)
ss.setdefault("last_audio", "")
ss.setdefault("n", 0)

if not ss.splash_done:
    ph = st.empty()
    ph.markdown(f'<div class="splash">{LOGO_SVG.format(s=180)}'
                '<p class="kora-title">KORA AI</p><p class="kora-sub">SYSTEM ONLINE</p></div>',
                unsafe_allow_html=True)
    time.sleep(1.8)
    ph.empty()
    ss.splash_done = True

# ---------- Funzioni ----------
def clean(t: str) -> str:
    return re.sub(r"[*_`#>]", "", t)

def make_pdf(text: str) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title="KORA AI")
    st_ = getSampleStyleSheet()
    story = [Paragraph("KORA AI", st_["Title"]), Spacer(1, 12)]
    for line in text.split("\n"):
        if line.strip():
            story.append(Paragraph(escape(clean(line)), st_["BodyText"]))
            story.append(Spacer(1, 6))
    doc.build(story)
    return buf.getvalue()

def make_docx(text: str) -> bytes:
    d = Document()
    d.add_heading("KORA AI", 0)
    for line in text.split("\n"):
        if line.startswith("#"):
            d.add_heading(line.lstrip("# ").strip(), 2)
        elif line.strip():
            d.add_paragraph(clean(line))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()

def gen_config(**kw):
    # thinking_budget=0 -> risposte molto più rapide
    return types.GenerateContentConfig(
        system_instruction=SYSTEM_INSTRUCTION, temperature=0.4,
        thinking_config=types.ThinkingConfig(thinking_budget=0), **kw)

def history_contents():
    out = []
    for m in ss.messages:
        if m["content"]:
            role = "user" if m["role"] == "user" else "model"
            out.append(types.Content(role=role, parts=[types.Part(text=m["content"])]))
    return out

def stream_answer():
    for chunk in client.models.generate_content_stream(
            model=MODEL, contents=history_contents(), config=gen_config()):
        if chunk.text:
            yield chunk.text

def transcribe(audio_bytes: bytes) -> str:
    r = client.models.generate_content(
        model=MODEL,
        contents=[types.Part.from_bytes(data=audio_bytes, mime_type="audio/wav"),
                  "Trascrivi fedelmente in italiano questo audio. Restituisci solo il testo."],
        config=types.GenerateContentConfig(
            thinking_config=types.ThinkingConfig(thinking_budget=0)))
    return (r.text or "").strip()

def speak(text: str):
    try:
        from gtts import gTTS
        short = clean(text)[:700]
        buf = io.BytesIO()
        gTTS(short, lang="it").write_to_fp(buf)
        st.audio(buf.getvalue(), format="audio/mp3", autoplay=True)
    except Exception as e:
        st.caption(f"Voce non disponibile: {e}")

def generate_image(prompt_it: str) -> str:
    """Traduce il prompt in inglese con Gemini e costruisce l'URL diretto Pollinations (gratis, senza chiave)."""
    prompt_in_inglese = client.models.generate_content(
        model=MODEL,
        contents="Riscrivi in inglese come prompt dettagliato per un generatore di immagini "
                 f"(solo il prompt, una riga, niente altro): {prompt_it}",
        config=types.GenerateContentConfig(
            thinking_config=types.ThinkingConfig(thinking_budget=0))).text.strip()
    seed = random.randint(1, 999999)
    return (f"https://image.pollinations.ai/prompt/{quote(prompt_in_inglese)}"
            f"?width=1024&height=1024&model=flux&nologo=true&seed={seed}")

def render_message(m, idx):
    with st.chat_message(m["role"], avatar="🧑‍💻" if m["role"] == "user" else "🤖"):
        if m.get("image_url"):
            st.image(m["image_url"], use_container_width=True)
            st.link_button("⬇️ Apri / scarica immagine", m["image_url"])
        if m["content"]:
            st.markdown(m["content"])
        if m["role"] == "assistant" and m["content"] and not m.get("image_url"):
            c1, c2, _ = st.columns([1, 1, 2])
            c1.download_button("📄 PDF", make_pdf(m["content"]), f"kora_{idx}.pdf",
                               "application/pdf", key=f"pdf{idx}")
            c2.download_button("📝 Word", make_docx(m["content"]), f"kora_{idx}.docx",
                               "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                               key=f"doc{idx}")

# ---------- UI ----------
c_logo, c_t = st.columns([1, 4])
c_logo.markdown(LOGO_SVG.format(s=64), unsafe_allow_html=True)
c_t.markdown('<p class="kora-title">KORA AI</p><p class="kora-sub">FUTURE ASSISTANT</p>',
             unsafe_allow_html=True)

with st.sidebar:
    st.markdown("### ⚙️ Controlli chiamata")
    mic_on = st.toggle("🎙️ Microfono attivo", value=False)
    speaker_on = st.toggle("🔊 Vivavoce (risposta a voce)", value=False)
    if st.button("🗑️ Nuova chat"):
        ss.messages = []
        st.rerun()

for i, m in enumerate(ss.messages):
    render_message(m, i)

prompt = st.chat_input("Scrivi a KORA...")
spoken = False
if mic_on:
    audio = st.audio_input("🎙️ Parla con KORA")
    if audio:
        data = audio.getvalue()
        h = hashlib.md5(data).hexdigest()
        if h != ss.last_audio:
            ss.last_audio = h
            with st.spinner("Ascolto..."):
                prompt = transcribe(data)
            spoken = True

if prompt:
    ss.messages.append({"role": "user", "content": prompt})
    render_message(ss.messages[-1], len(ss.messages) - 1)
    try:
        if IMG_RE.search(prompt):
            with st.chat_message("assistant", avatar="🤖"):
                with st.spinner("Creo l'immagine..."):
                    url_immagine = generate_image(prompt)
                st.image(url_immagine, use_container_width=True)
            ss.messages.append({"role": "assistant", "content": "Ecco la tua immagine.",
                                "image_url": url_immagine})
        else:
            with st.chat_message("assistant", avatar="🤖"):
                text = st.write_stream(stream_answer())
                if speaker_on or spoken:
                    speak(text)
            ss.messages.append({"role": "assistant", "content": text})
    except Exception as e:
        st.error(f"Errore: {e}")
    if ss.messages[-1]["role"] == "assistant" and not (speaker_on or spoken):
        st.rerun()  # mostra subito i pulsanti PDF/Word
